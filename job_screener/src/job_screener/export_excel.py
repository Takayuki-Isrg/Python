"""採点結果を Excel に出力する（openpyxl で色分け・ハイパーリンク・フィルタ）。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Evaluation, Job

# ランクごとの行の色
_RANK_FILL = {
    "A": "C6EFCE",   # 緑
    "B": "E2EFDA",   # 薄緑
    "C": "F2F2F2",   # 灰
    "D": "FFC7CE",   # 赤
    "?": "FFEB9C",   # 橙
}
_RANK_ORDER = {"A": 0, "B": 1, "?": 2, "C": 3, "D": 4}

_HEADERS = [
    ("ランク", 7), ("スコア", 7), ("会社名", 28), ("求人名", 60),
    ("年収下限", 10), ("年収上限", 10), ("プラス要素", 30),
    ("懸念点", 30), ("要確認事項", 34), ("既読", 6), ("経過日", 7), ("求人URL", 22),
]


def _man(v: int | None) -> str:
    return f"{v // 10000}万" if v else ""


def export(results: list[tuple[Job, Evaluation]], out_path: str | Path) -> Path:
    # ランク順 → スコア降順で並べる（人間が上から読む順）
    results = sorted(results, key=lambda re: (_RANK_ORDER.get(re[1].rank, 9), -re[1].score))

    wb = Workbook()
    ws = wb.active
    ws.title = "求人選別"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for col, (name, width) in enumerate(_HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    for job, ev in results:
        row = [
            ev.rank, ev.score, job.company, job.title,
            _man(job.salary_min), _man(job.salary_max),
            " / ".join(ev.pluses), " / ".join(ev.concerns),
            " / ".join(ev.unknowns),
            "既読" if job.is_read else "未読",
            job.days_since, "リンク",
        ]
        ws.append(row)
        r = ws.max_row
        fill = PatternFill("solid", fgColor=_RANK_FILL.get(ev.rank, "FFFFFF"))
        for col in range(1, len(_HEADERS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=col in (4, 7, 8, 9))
        # URL をハイパーリンク化
        link_cell = ws.cell(row=r, column=len(_HEADERS))
        if job.url:
            link_cell.hyperlink = job.url
            link_cell.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_HEADERS))}{ws.max_row}"

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
